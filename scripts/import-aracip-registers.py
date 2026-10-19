#!/usr/bin/env python3
"""Import the current ARACIP registers into a SIIIR-keyed JSON layer."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCHOOLS_PATH = ROOT / "src" / "data" / "schools.json"
OUTPUT_PATH = ROOT / "src" / "data" / "aracip-status.json"
SOURCE_URL = "https://aracip.eu/categorii-documente/info-unitati-invatamant-registre"
COUNTY_CODES = {
    "ALBA": "AB", "ARAD": "AR", "ARGES": "AG", "BACAU": "BC", "BIHOR": "BH",
    "BISTRITA NASAUD": "BN", "BOTOSANI": "BT", "BRASOV": "BV", "BRAILA": "BR",
    "BUZAU": "BZ", "CARAS SEVERIN": "CS", "CALARASI": "CL", "CLUJ": "CJ",
    "CONSTANTA": "CT", "COVASNA": "CV", "DAMBO VITA": "DB", "DAMBOVITA": "DB",
    "DOLJ": "DJ", "GALATI": "GL", "GIURGIU": "GR", "GORJ": "GJ", "HARGHITA": "HR",
    "HUNEDOARA": "HD", "IALOMITA": "IL", "IASI": "IS", "ILFOV": "IF", "MARAMURES": "MM",
    "MEHEDINTI": "MH", "MURES": "MS", "NEAMT": "NT", "OLT": "OT", "PRAHOVA": "PH",
    "SALAJ": "SJ", "SATU MARE": "SM", "SIBIU": "SB", "SUCEAVA": "SV", "TELEORMAN": "TR",
    "TIMIS": "TM", "TULCEA": "TL", "VASLUI": "VS", "VALCEA": "VL", "VRANCEA": "VN",
    "BUCURESTI": "B", "MUNICIPIUL BUCURESTI": "B",
}


def norm(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value or "").lower())
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def tokens(value: object) -> set[str]:
    return {word for word in norm(value).split() if len(word) > 2}


def text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()


def row_text(row: tuple[object, ...]) -> str:
    return " ".join(text(value) for value in row if value not in (None, ""))


def parse_workbook(path: Path, register_type: str, updated_at: str) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        # A few ARACIP sheets carry Excel's full 16,384-column dimension even
        # though their data lives in the first handful of columns. Limiting
        # the read keeps the import fast and avoids materialising millions of
        # empty cells.
        rows = list(sheet.iter_rows(values_only=True, max_col=min(sheet.max_column or 32, 32)))
        header_index = None
        header = []
        for index, row in enumerate(rows):
            values = [norm(value) for value in row]
            if any("denumirea unitatii" in value for value in values):
                header_index = index
                header = values
                break
        if header_index is None:
            continue
        name_index = next((i for i, value in enumerate(header) if "denumirea unitatii" in value), 1)
        county_index = next((i for i, value in enumerate(header) if value in {"judet", "judetul"} or value.startswith("judet ")), None)
        county = ""
        # Most register sheets print the county heading immediately before the
        # column header, so capture it before iterating the data rows.
        for previous in reversed(rows[:header_index]):
            match = re.search(r"JUDETUL\s*:?\s*([A-Z][A-Z -]+)", norm(row_text(previous)).upper())
            if match:
                county = COUNTY_CODES.get(match.group(1).strip(), match.group(1).strip()[:2])
                break
        for row in rows[header_index + 1:]:
            joined = row_text(row)
            county_match = re.search(r"JUDETUL\s*:?\s*([A-Z][A-Z -]+)", norm(joined).upper())
            if county_match:
                county = COUNTY_CODES.get(county_match.group(1).strip(), county_match.group(1).strip()[:2])
                continue
            name = text(row[name_index] if name_index < len(row) else "")
            if not name or not re.search(r"[A-Za-zĂÂÎȘȚăâîșț]{3}", name):
                continue
            if norm(name).startswith(("denumirea unitatii", "persoanei juridice", "unitatea de invatamant")):
                continue
            # Data rows have a numbered first column in all current ARACIP books.
            first = text(row[0] if row else "")
            if not re.match(r"^\d+(?:\.0)?$", first):
                continue
            row_county = text(row[county_index] if county_index is not None and county_index < len(row) else "")
            if row_county and re.fullmatch(r"[A-Za-z]{2}", row_county):
                county = row_county.upper()
            contact_index = next((i for i, value in enumerate(header) if "date de contact" in value), None)
            next_eval_index = next((i for i, value in enumerate(header) if "urmatoarea evaluare" in value), None)
            order_index = next((i for i, value in enumerate(header) if "ordin" in value), None)
            evaluation_year_index = next((i for i, value in enumerate(header) if "anul evaluarii" in value), None)
            records.append({
                "name": name,
                "county": county,
                "contact": text(row[contact_index]) if contact_index is not None and contact_index < len(row) else "",
                "order": text(row[order_index]) if order_index is not None and order_index < len(row) else "",
                "evaluationYear": text(row[evaluation_year_index]) if evaluation_year_index is not None and evaluation_year_index < len(row) else "",
                "nextEvaluationSchoolYear": text(row[next_eval_index]) if next_eval_index is not None and next_eval_index < len(row) else "",
                "registerType": register_type,
                "updatedAt": updated_at,
                "source": SOURCE_URL,
            })
    return records


def build_school_indexes(schools: list[dict[str, object]]) -> tuple[dict[str, list[dict[str, object]]], dict[str, dict[str, object]]]:
    by_token: dict[str, list[dict[str, object]]] = {}
    by_name: dict[str, dict[str, object]] = {}
    for school in schools:
        by_name.setdefault(norm(school["name"]), school)
        for token in tokens(school["name"]):
            by_token.setdefault(token, []).append(school)
    return by_token, by_name


def find_school(record: dict[str, str], schools: list[dict[str, object]], by_token: dict[str, list[dict[str, object]]], by_name: dict[str, dict[str, object]]) -> dict[str, object] | None:
    name_tokens = tokens(record["name"])
    if not name_tokens:
        return None
    county = norm(record.get("county"))
    exact = by_name.get(norm(record["name"]))
    if exact and (not county or norm(exact.get("county")) == county):
        return exact
    # Restrict matching to schools sharing at least one meaningful token. This
    # avoids the quadratic 18k x ARACIP-record comparison while retaining a
    # fuzzy fallback for renamed units and names containing multiple aliases.
    candidate_ids: dict[str, dict[str, object]] = {}
    # Start with the rarest name tokens; common tokens such as "scoala" can
    # otherwise expand the candidate set to thousands of rows.
    indexed_tokens = sorted(name_tokens, key=lambda token: len(by_token.get(token, [])))[:3]
    for token in indexed_tokens:
        for school in by_token.get(token, []):
            if county and norm(school.get("county")) != county:
                continue
            candidate_ids[str(school["siiirCode"])] = school
    candidates = list(candidate_ids.values())
    best: tuple[float, dict[str, object]] | None = None
    for school in candidates:
        school_tokens = tokens(school["name"])
        overlap = len(name_tokens & school_tokens) / max(1, len(name_tokens | school_tokens))
        containment = name_tokens.issubset(school_tokens) or school_tokens.issubset(name_tokens)
        score = overlap + (0.2 if containment else 0)
        if best is None or score > best[0]:
            best = (score, school)
    return best[1] if best and best[0] >= 0.58 else None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    args = parser.parse_args()
    schools = json.loads(SCHOOLS_PATH.read_text(encoding="utf-8"))
    by_token, by_name = build_school_indexes(schools)
    files = [
        (args.directory / "acreditate.xlsx", "acreditata", "2026-10-06"),
        (args.directory / "privat-provizoriu.xlsx", "autorizata_provizoriu_particular", "2026-09-10"),
        (args.directory / "stat-provizoriu.xlsx", "autorizata_provizoriu_stat", "2026-09-10"),
        (args.directory / "evaluare-periodica.xlsx", "evaluata_periodic", "2026-09-10"),
    ]
    merged: dict[str, dict[str, object]] = {}
    matched = 0
    for path, register_type, updated_at in files:
        if not path.exists():
            continue
        for record in parse_workbook(path, register_type, updated_at):
            school = find_school(record, schools, by_token, by_name)
            if not school:
                continue
            matched += 1
            code = str(school["siiirCode"])
            current = merged.setdefault(code, {
                "siiirCode": code,
                "schoolName": school["name"],
                "county": school["county"],
                "source": SOURCE_URL,
                "sources": [],
            })
            current["sources"].append({
                "registerType": record["registerType"],
                "updatedAt": record["updatedAt"],
                "contact": record["contact"],
                "order": record["order"],
                "evaluationYear": record.get("evaluationYear", ""),
                "nextEvaluationSchoolYear": record["nextEvaluationSchoolYear"],
            })
            if record["registerType"] == "evaluata_periodic":
                if "registerType" not in current:
                    current["registerType"] = "evaluata_periodic"
                current["lastExternalEvaluation"] = record.get("evaluationYear") or None
                current["nextEvaluationSchoolYear"] = record["nextEvaluationSchoolYear"]
            elif "registerType" not in current or current["registerType"] != "acreditata":
                current["registerType"] = record["registerType"]
                current["accreditationOrder"] = record["order"]
    OUTPUT_PATH.write_text(json.dumps(merged, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Înregistrări ARACIP asociate: {matched}; școli unice: {len(merged)}")
    print(f"Output: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
